"""
ETL Pipeline for PPCAC (Programa de Prevenção de Conflitos Agrários Coletivos).
Ingests data/raw/Planilha_da_Relacao_Conflitos_09.02.2026.xlsx into PostGIS table
public.ppcac_conflitos_pe with tiered spatial resolution (DataJud -> Municipal fallback).
"""

import logging
import os
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import openpyxl
from sqlalchemy import text
from src.database import get_engine

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

DEFAULT_SPREADSHEET = BASE_DIR / "data" / "raw" / "Planilha_da_Relacao_Conflitos_09.02.2026.xlsx"


def _norm(s: Optional[str]) -> str:
    if not s:
        return ""
    normalized = unicodedata.normalize("NFKD", str(s)).encode("ASCII", "ignore").decode("ASCII")
    return re.sub(r"[^a-z0-9 ]", " ", normalized.lower()).strip()


def _clean_area_name(name: str) -> str:
    if not name:
        return ""
    cleaned = re.sub(
        r"^(?:engenhos?|fazendas?|sitios?|glebas?|comunidades?\s+quilombolas?|comunidades?|assentamentos?|acampamentos?|povoados?|loteamentos?|nucleos?|terrenos?|imoveis?)\s+",
        "",
        name,
        flags=re.I
    ).strip()
    return cleaned


def _is_cnj(val: str) -> bool:
    digits = re.sub(r"\D", "", val)
    if len(digits) in (19, 20) or (".8.17." in val or ".4.05." in val or "-8300" in val or "-2940" in val or "numero errado" in val):
        return True
    return False


def _is_mppe(val: str) -> bool:
    s = val.strip()
    if re.search(r"02\d{3}\.\d{3}\.\d{3}/\d{4}", s) or s.startswith("0205") or s.startswith("0207") or s.startswith("1.26.000"):
        return True
    return False


def _is_sei(val: str) -> bool:
    s = val.strip()
    if re.search(r"\d{10}\.\d{6}/\d{4}-\d{2}", s):
        return True
    if re.match(r"^(?:OF|OF\.|370|390|220|003|190|250)\b", s):
        return True
    return False


def parse_ppcac_spreadsheet(filepath: Path | str = DEFAULT_SPREADSHEET) -> List[Dict[str, Any]]:
    """Parse Excel spreadsheet, forward-fill merged cells, rectify shifted columns,
    and consolidate into canonical conflict areas."""
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"Spreadsheet not found at {path}")

    wb = openpyxl.load_workbook(str(path), data_only=True)
    sheet = wb["Planilha1"]

    curr_area: Optional[str] = None
    curr_mun: Optional[str] = None

    raw_records: List[Dict[str, Any]] = []

    for r in range(2, sheet.max_row + 1):
        cell_area = sheet.cell(r, 1).value
        cell_mun = sheet.cell(r, 2).value

        # Check if entire row is empty
        row_vals = [sheet.cell(r, c).value for c in range(1, 9)]
        if not any(v is not None and str(v).strip() != "" for v in row_vals):
            continue

        if cell_area is not None and str(cell_area).strip() != "":
            curr_area = str(cell_area).strip()
            curr_mun = str(cell_mun).strip() if cell_mun else None
        else:
            # Merged cell forward-fill
            if curr_area is None:
                continue

        # Examine columns 3 to 8
        proprietario: Optional[str] = None
        movimento: Optional[str] = None
        judiciais: List[str] = []
        mppe: List[str] = []
        sei: List[str] = []
        ano_ref: Optional[int] = None
        situacao = "ATIVO"
        observacoes: List[str] = []

        for c_idx in range(3, 9):
            val = sheet.cell(r, c_idx).value
            if val is None:
                continue
            val_str = str(val).strip()
            if not val_str:
                continue

            # Year number or float (e.g. 2026.0)
            if isinstance(val, (int, float)) and 2000 <= val <= 2030:
                ano_ref = int(val)
                continue

            # Check status
            if "ARQUIVADO" in val_str.upper():
                situacao = "ARQUIVADO"
                continue

            # Split multiline entries
            lines = [line.strip() for line in val_str.split("\n") if line.strip()]

            for line in lines:
                if "ARQUIVADO" in line.upper():
                    situacao = "ARQUIVADO"
                    continue
                if re.match(r"^\d{4}(?:\.0)?$", line) and 2000 <= float(line) <= 2030:
                    ano_ref = int(float(line))
                    continue

                if _is_mppe(line):
                    mppe.append(line)
                elif _is_sei(line):
                    sei.append(line)
                elif _is_cnj(line):
                    # Clean any trailing note like "numero errado"
                    cleaned_cnj = re.sub(r"\s+numero errado.*$", "", line, flags=re.IGNORECASE).strip()
                    judiciais.append(cleaned_cnj)
                else:
                    # Could be owner or movement based on column or contents
                    if c_idx == 3:
                        proprietario = line
                    elif c_idx == 4:
                        movimento = line
                    else:
                        observacoes.append(line)

        raw_records.append({
            "nome_area": curr_area,
            "municipio": curr_mun or "",
            "proprietario": proprietario,
            "movimento_social": movimento,
            "processos_judiciais": judiciais,
            "processos_mppe": mppe,
            "processos_sei": sei,
            "ano_referencia": ano_ref,
            "situacao": situacao,
            "observacoes": observacoes,
        })

    # Group into canonical unique areas by nome_area
    areas_dict: Dict[str, Dict[str, Any]] = {}

    for rec in raw_records:
        area_name = rec["nome_area"]
        # Handle known typo in raw spreadsheet where Jaqueira was typed as BARRO BRANCO
        mun_clean = rec["municipio"]
        if mun_clean.upper() == "BARRO BRANCO":
            mun_clean = "Jaqueira"

        if area_name not in areas_dict:
            areas_dict[area_name] = {
                "nome_area": area_name,
                "municipio": mun_clean,
                "proprietario": rec["proprietario"],
                "movimento_social": rec["movimento_social"],
                "processos_judiciais": list(rec["processos_judiciais"]),
                "processos_mppe": list(rec["processos_mppe"]),
                "processos_sei": list(rec["processos_sei"]),
                "ano_referencia": rec["ano_referencia"],
                "situacao": rec["situacao"],
                "observacoes": list(rec["observacoes"]),
            }
        else:
            curr = areas_dict[area_name]
            if (not curr["municipio"] or curr["municipio"].upper() == "BARRO BRANCO") and mun_clean:
                curr["municipio"] = mun_clean
            if not curr["proprietario"] and rec["proprietario"]:
                curr["proprietario"] = rec["proprietario"]
            if not curr["movimento_social"] and rec["movimento_social"]:
                curr["movimento_social"] = rec["movimento_social"]
            if not curr["ano_referencia"] and rec["ano_referencia"]:
                curr["ano_referencia"] = rec["ano_referencia"]
            if rec["situacao"] == "ARQUIVADO":
                curr["situacao"] = "ARQUIVADO"

            for j in rec["processos_judiciais"]:
                if j not in curr["processos_judiciais"]:
                    curr["processos_judiciais"].append(j)
            for m in rec["processos_mppe"]:
                if m not in curr["processos_mppe"]:
                    curr["processos_mppe"].append(m)
            for s in rec["processos_sei"]:
                if s not in curr["processos_sei"]:
                    curr["processos_sei"].append(s)
            for o in rec["observacoes"]:
                if o not in curr["observacoes"]:
                    curr["observacoes"].append(o)

    # Convert to list
    consolidated = list(areas_dict.values())
    logger.info(f"Consolidated {len(consolidated)} unique conflict areas from spreadsheet.")
    return consolidated


def ingest_ppcac(engine=None) -> int:
    """Ingest consolidated PPCAC areas into PostGIS table public.ppcac_conflitos_pe."""
    if engine is None:
        engine = get_engine()

    areas = parse_ppcac_spreadsheet()

    with engine.begin() as conn:
        # Load municipal lookup
        mun_lookup: Dict[str, Tuple[int, str, float, float, List[float]]] = {}
        res = conn.execute(text("""
            SELECT municipio_ibge, municipio_nome, 
                   ST_X(ST_Centroid(geometry)) as lon, 
                   ST_Y(ST_Centroid(geometry)) as lat,
                   ST_XMin(geometry) as minx, ST_YMin(geometry) as miny,
                   ST_XMax(geometry) as maxx, ST_YMax(geometry) as maxy
            FROM public.jurisdicoes_pe_municipios
        """)).fetchall()
        for r in res:
            # normalized name -> (ibge, name, lon, lat, bbox)
            mun_lookup[_norm(r[1])] = (r[0], r[1], r[2], r[3], [r[4], r[5], r[6], r[7]])

        # Load DataJud lookup
        dj_lookup: Dict[str, Tuple[float, float, str]] = {}
        dj_res = conn.execute(text("""
            SELECT regexp_replace(numero_processo, '\\D', '', 'g') as digits,
                   numero_processo,
                   ST_X(geometry) as lon,
                   ST_Y(geometry) as lat
            FROM public.processos_conflitos_judiciais
            WHERE geometry IS NOT NULL
        """)).fetchall()
        for r in dj_res:
            dj_lookup[r[0]] = (r[2], r[3], r[1])

        # Load Despejo Zero lookup
        dz_lookup: List[Tuple[str, str, float, float]] = []
        dz_res = conn.execute(text("""
            SELECT conflito_id, nome_comunidade,
                   ST_X(geometry) as lon, ST_Y(geometry) as lat
            FROM public.despejo_zero_pe
            WHERE geometry IS NOT NULL
        """)).fetchall()
        for r in dz_res:
            dz_lookup.append((r[0], r[1], r[2], r[3]))

        # Recreate table
        conn.execute(text("DROP TABLE IF EXISTS public.ppcac_conflitos_pe CASCADE;"))
        conn.execute(text("""
            CREATE TABLE public.ppcac_conflitos_pe (
                id SERIAL PRIMARY KEY,
                programa TEXT NOT NULL DEFAULT 'PPCAC',
                nome_area TEXT NOT NULL,
                municipio TEXT NOT NULL,
                municipio_ibge INTEGER REFERENCES public.jurisdicoes_pe_municipios(municipio_ibge),
                proprietario TEXT,
                movimento_social TEXT,
                processos_judiciais TEXT[] DEFAULT '{}',
                processos_mppe TEXT[] DEFAULT '{}',
                processos_sei TEXT[] DEFAULT '{}',
                ano_referencia INTEGER,
                situacao TEXT NOT NULL DEFAULT 'ATIVO',
                observacoes TEXT,
                datajud_processos TEXT[] DEFAULT '{}',
                despejo_zero_ids TEXT[] DEFAULT '{}',
                sigef_codigos TEXT[] DEFAULT '{}',
                car_codigos TEXT[] DEFAULT '{}',
                iterpe_nomes TEXT[] DEFAULT '{}',
                incra_projetos TEXT[] DEFAULT '{}',
                total_car_imoveis INTEGER DEFAULT 0,
                fonte_geometria TEXT NOT NULL DEFAULT 'MUNICIPIO',
                tipo_geometria TEXT NOT NULL DEFAULT 'POINT',
                tem_geometria_exata BOOLEAN NOT NULL DEFAULT FALSE,
                centroid_lat DOUBLE PRECISION,
                centroid_lon DOUBLE PRECISION,
                bbox DOUBLE PRECISION[],
                geometry GEOMETRY(Geometry, 4326),
                created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
            );
        """))

        # Prepare records for insertion
        inserted_count = 0
        for a in areas:
            norm_m = _norm(a["municipio"])
            ibge_info = mun_lookup.get(norm_m)
            if not ibge_info:
                # Fuzzy match municipality
                for k, v in mun_lookup.items():
                    if k in norm_m or norm_m in k:
                        ibge_info = v
                        break

            municipio_ibge = ibge_info[0] if ibge_info else None
            mun_name = ibge_info[1] if ibge_info else a["municipio"]
            mun_lon = ibge_info[2] if ibge_info else -37.0
            mun_lat = ibge_info[3] if ibge_info else -8.0
            mun_bbox = ibge_info[4] if ibge_info else [-37.5, -8.5, -36.5, -7.5]

            p_name = a["nome_area"]
            core_area = _clean_area_name(p_name)
            norm_core = _norm(core_area)

            geom_wkt = None
            fonte_geometria = "MUNICIPIO"
            tipo_geometria = "POINT"
            tem_geometria_exata = False
            sigef_matched_codes: List[str] = []
            car_matched_codes: List[str] = []
            iterpe_matched_names: List[str] = []
            incra_matched_projs: List[str] = []
            centroid_lat = mun_lat
            centroid_lon = mun_lon
            bbox = mun_bbox

            # Tier 1: Match sigef_casos_analisados (e.g. Engenho Batateiras in Maraial)
            ca_match = conn.execute(text("""
                SELECT nome_imovel,
                       ST_AsText(ST_Force2D(ST_Union(geometry))) as wkt,
                       ST_XMin(ST_Union(geometry)) as minx, ST_YMin(ST_Union(geometry)) as miny,
                       ST_XMax(ST_Union(geometry)) as maxx, ST_YMax(ST_Union(geometry)) as maxy,
                       ST_X(ST_Centroid(ST_Union(geometry))) as cx, ST_Y(ST_Centroid(ST_Union(geometry))) as cy
                FROM public.sigef_casos_analisados
                WHERE lower(unaccent(nome_imovel)) LIKE :p
                GROUP BY nome_imovel
                LIMIT 1
            """), {"p": f"%{norm_core}%"}).mappings().first()

            if ca_match:
                fonte_geometria = "SIGEF_CASOS_ANALISADOS"
                tipo_geometria = "POLYGON"
                tem_geometria_exata = True
                sigef_matched_codes.append(str(ca_match["nome_imovel"]))
                geom_wkt = ca_match["wkt"]
                bbox = [float(ca_match["minx"]), float(ca_match["miny"]), float(ca_match["maxx"]), float(ca_match["maxy"])]
                centroid_lon = float(ca_match["cx"])
                centroid_lat = float(ca_match["cy"])

            # Tier 2a: Match sigef_privado_pe in the same municipality (e.g. Roncadorzinho, Fervedouro, etc.)
            if not geom_wkt and len(core_area) >= 4 and municipio_ibge:
                sp_match = conn.execute(text("""
                    SELECT codigo_imo, nome_area,
                           ST_AsText(ST_Force2D(ST_Union(geometry))) as wkt,
                           ST_XMin(ST_Union(geometry)) as minx, ST_YMin(ST_Union(geometry)) as miny,
                           ST_XMax(ST_Union(geometry)) as maxx, ST_YMax(ST_Union(geometry)) as maxy,
                           ST_X(ST_Centroid(ST_Union(geometry))) as cx, ST_Y(ST_Centroid(ST_Union(geometry))) as cy
                    FROM public.sigef_privado_pe
                    WHERE municipio = :ibge
                      AND lower(unaccent(nome_area)) ~* ('(^|[^a-z0-9])' || :regex_core || '([^a-z0-9]|$)')
                      AND similarity(lower(unaccent(nome_area)), :full) >= 0.35
                    GROUP BY codigo_imo, nome_area
                    ORDER BY similarity(lower(unaccent(nome_area)), :full) DESC
                    LIMIT 1
                """), {"ibge": municipio_ibge, "regex_core": re.escape(norm_core), "full": _norm(p_name)}).mappings().first()

                if sp_match:
                    fonte_geometria = "SIGEF_PRIVADO"
                    tipo_geometria = "POLYGON"
                    tem_geometria_exata = True
                    sigef_matched_codes.append(str(sp_match["codigo_imo"]))
                    geom_wkt = sp_match["wkt"]
                    bbox = [float(sp_match["minx"]), float(sp_match["miny"]), float(sp_match["maxx"]), float(sp_match["maxy"])]
                    centroid_lon = float(sp_match["cx"])
                    centroid_lat = float(sp_match["cy"])

            # Tier 2b: Match sigef_publico_pe in the same municipality (e.g. PA São Pedro in Jaboatão)
            if not geom_wkt and len(core_area) >= 4 and municipio_ibge:
                spub_match = conn.execute(text("""
                    SELECT codigo_imo, nome_area,
                           ST_AsText(ST_Force2D(ST_Union(geometry))) as wkt,
                           ST_XMin(ST_Union(geometry)) as minx, ST_YMin(ST_Union(geometry)) as miny,
                           ST_XMax(ST_Union(geometry)) as maxx, ST_YMax(ST_Union(geometry)) as maxy,
                           ST_X(ST_Centroid(ST_Union(geometry))) as cx, ST_Y(ST_Centroid(ST_Union(geometry))) as cy
                    FROM public.sigef_publico_pe
                    WHERE municipio = :ibge
                      AND lower(unaccent(nome_area)) ~* ('(^|[^a-z0-9])' || :regex_core || '([^a-z0-9]|$)')
                      AND similarity(lower(unaccent(nome_area)), :full) >= 0.35
                    GROUP BY codigo_imo, nome_area
                    ORDER BY similarity(lower(unaccent(nome_area)), :full) DESC
                    LIMIT 1
                """), {"ibge": municipio_ibge, "regex_core": re.escape(norm_core), "full": _norm(p_name)}).mappings().first()

                if spub_match:
                    fonte_geometria = "SIGEF_PUBLICO"
                    tipo_geometria = "POLYGON"
                    tem_geometria_exata = True
                    sigef_matched_codes.append(str(spub_match["codigo_imo"]))
                    geom_wkt = spub_match["wkt"]
                    bbox = [float(spub_match["minx"]), float(spub_match["miny"]), float(spub_match["maxx"]), float(spub_match["maxy"])]
                    centroid_lon = float(spub_match["cx"])
                    centroid_lat = float(spub_match["cy"])

            # Tier 2c: Match iterpe_glebas_pe (e.g. Comunidade Quilombola Negros de Gilú in Itacuruba)
            if not geom_wkt and len(core_area) >= 4:
                iterpe_match = conn.execute(text("""
                    SELECT id, nome, tipo,
                           ST_AsText(ST_Force2D(ST_Union(geometry))) as wkt,
                           ST_XMin(ST_Union(geometry)) as minx, ST_YMin(ST_Union(geometry)) as miny,
                           ST_XMax(ST_Union(geometry)) as maxx, ST_YMax(ST_Union(geometry)) as maxy,
                           ST_X(ST_Centroid(ST_Union(geometry))) as cx, ST_Y(ST_Centroid(ST_Union(geometry))) as cy
                    FROM public.iterpe_glebas_pe
                    WHERE lower(unaccent(nome)) ~* ('(^|[^a-z0-9])' || :regex_core || '([^a-z0-9]|$)')
                      AND similarity(lower(unaccent(nome)), :full) >= 0.35
                      AND (lower(unaccent(municipio)) = :mun OR :mun = '')
                    GROUP BY id, nome, tipo
                    LIMIT 1
                """), {"regex_core": re.escape(norm_core), "full": _norm(p_name), "mun": norm_m}).mappings().first()

                if iterpe_match:
                    fonte_geometria = "ITERPE"
                    tipo_geometria = "POLYGON"
                    tem_geometria_exata = True
                    iterpe_matched_names.append(f"{iterpe_match['nome']} ({iterpe_match['tipo']})")
                    geom_wkt = iterpe_match["wkt"]
                    bbox = [float(iterpe_match["minx"]), float(iterpe_match["miny"]), float(iterpe_match["maxx"]), float(iterpe_match["maxy"])]
                    centroid_lon = float(iterpe_match["cx"])
                    centroid_lat = float(iterpe_match["cy"])

            # Tier 2d: Match assentamentos_incra_pe (e.g. PA São Gregório/Alegre in Gameleira)
            if not geom_wkt and len(core_area) >= 4:
                parts = [_clean_area_name(p) for p in re.split(r"[,;/e]\s*", p_name) if len(_clean_area_name(p)) >= 4]
                for part in parts:
                    incra_match = conn.execute(text("""
                        SELECT id, no_projeto,
                               ST_AsText(ST_Force2D(ST_Union(geometry))) as wkt,
                               ST_XMin(ST_Union(geometry)) as minx, ST_YMin(ST_Union(geometry)) as miny,
                               ST_XMax(ST_Union(geometry)) as maxx, ST_YMax(ST_Union(geometry)) as maxy,
                               ST_X(ST_Centroid(ST_Union(geometry))) as cx, ST_Y(ST_Centroid(ST_Union(geometry))) as cy
                        FROM public.assentamentos_incra_pe
                        WHERE lower(unaccent(no_projeto)) ~* ('(^|[^a-z0-9])' || :regex_part || '([^a-z0-9]|$)')
                          AND (lower(unaccent(no_municipio)) = :mun OR :mun = '')
                        GROUP BY id, no_projeto
                        LIMIT 1
                    """), {"regex_part": re.escape(_norm(part)), "mun": norm_m}).mappings().first()

                    if incra_match:
                        fonte_geometria = "INCRA"
                        tipo_geometria = "POLYGON"
                        tem_geometria_exata = True
                        incra_matched_projs.append(str(incra_match["no_projeto"]))
                        geom_wkt = incra_match["wkt"]
                        bbox = [float(incra_match["minx"]), float(incra_match["miny"]), float(incra_match["maxx"]), float(incra_match["maxy"])]
                        centroid_lon = float(incra_match["cx"])
                        centroid_lat = float(incra_match["cy"])
                        break

            # Match DataJud judicial lawsuit numbers
            matched_dj_processes = []
            exact_dj_coords = None
            for p in a["processos_judiciais"]:
                p_digits = re.sub(r"\D", "", p)
                if p_digits in dj_lookup:
                    lon, lat, orig_num = dj_lookup[p_digits]
                    matched_dj_processes.append(orig_num)
                    if not exact_dj_coords:
                        exact_dj_coords = (lon, lat)

            # Match Despejo Zero (using normalized names and municipality)
            matched_dz_ids = []
            exact_dz_coords = None
            for dz_id, dz_name, dz_lon, dz_lat in dz_lookup:
                norm_dz = _norm(dz_name)
                clean_dz = _clean_area_name(dz_name)
                if len(norm_core) >= 4 and (norm_core in norm_dz or _norm(clean_dz) in norm_core or norm_dz in norm_core):
                    matched_dz_ids.append(dz_id)
                    if not exact_dz_coords:
                        exact_dz_coords = (dz_lon, dz_lat)

            # Tier 3: DataJud fallback point (if no polygon)
            if not geom_wkt and exact_dj_coords:
                fonte_geometria = "DATAJUD"
                tipo_geometria = "POINT"
                tem_geometria_exata = True
                centroid_lon, centroid_lat = exact_dj_coords
                bbox = [centroid_lon - 0.02, centroid_lat - 0.02, centroid_lon + 0.02, centroid_lat + 0.02]
                geom_wkt = f"POINT({centroid_lon} {centroid_lat})"

            # Tier 4: Despejo Zero fallback point
            elif not geom_wkt and exact_dz_coords:
                fonte_geometria = "DESPEJO_ZERO"
                tipo_geometria = "POINT"
                tem_geometria_exata = True
                centroid_lon, centroid_lat = exact_dz_coords
                bbox = [centroid_lon - 0.02, centroid_lat - 0.02, centroid_lon + 0.02, centroid_lat + 0.02]
                geom_wkt = f"POINT({centroid_lon} {centroid_lat})"

            # Tier 5: Municipal fallback point
            elif not geom_wkt:
                fonte_geometria = "MUNICIPIO"
                tipo_geometria = "POINT"
                tem_geometria_exata = False
                centroid_lon, centroid_lat = mun_lon, mun_lat
                bbox = mun_bbox
                geom_wkt = f"POINT({centroid_lon} {centroid_lat})"

            # Cross-referencing: Associate CAR parcels from area_imovel_1
            if geom_wkt:
                if tipo_geometria == "POLYGON":
                    car_res = conn.execute(text("""
                        SELECT DISTINCT cod_imovel
                        FROM public.area_imovel_1
                        WHERE ST_Intersects(geometry, ST_GeomFromText(:wkt, 4326))
                        LIMIT 50;
                    """), {"wkt": geom_wkt}).fetchall()
                    car_matched_codes.extend([r[0] for r in car_res if r[0]])
                elif tipo_geometria == "POINT" and fonte_geometria != "MUNICIPIO":
                    car_res = conn.execute(text("""
                        SELECT DISTINCT cod_imovel
                        FROM public.area_imovel_1
                        WHERE ST_Contains(geometry, ST_GeomFromText(:wkt, 4326))
                        LIMIT 10;
                    """), {"wkt": geom_wkt}).fetchall()
                    car_matched_codes.extend([r[0] for r in car_res if r[0]])

            # Cross-referencing: CAR Casos Analisados (by property name & municipality)
            if len(norm_core) >= 4:
                car_ca = conn.execute(text("""
                    SELECT cod_imovel
                    FROM public.car_casos_analisados
                    WHERE lower(unaccent(nome_imovel)) LIKE :p
                      AND (lower(unaccent(municipio)) = :mun OR :mun = '')
                """), {"p": f"%{norm_core}%", "mun": norm_m}).fetchall()
                for c_row in car_ca:
                    if c_row[0] and c_row[0] not in car_matched_codes:
                        car_matched_codes.append(str(c_row[0]))

            # Cross-referencing: Territorial overlaps with ITERPE and INCRA for polygon areas
            if tipo_geometria == "POLYGON" and geom_wkt:
                ov_iterpe = conn.execute(text("""
                    SELECT DISTINCT nome, tipo
                    FROM public.iterpe_glebas_pe
                    WHERE ST_Intersects(geometry, ST_GeomFromText(:wkt, 4326));
                """), {"wkt": geom_wkt}).fetchall()
                for oi in ov_iterpe:
                    entry = f"{oi[0]} ({oi[1]})"
                    if entry not in iterpe_matched_names:
                        iterpe_matched_names.append(entry)

                ov_incra = conn.execute(text("""
                    SELECT DISTINCT no_projeto
                    FROM public.assentamentos_incra_pe
                    WHERE ST_Intersects(geometry, ST_GeomFromText(:wkt, 4326));
                """), {"wkt": geom_wkt}).fetchall()
                for oinc in ov_incra:
                    if oinc[0] and oinc[0] not in incra_matched_projs:
                        incra_matched_projs.append(str(oinc[0]))

            obs_str = " | ".join(a["observacoes"]) if a["observacoes"] else None

            conn.execute(text("""
                INSERT INTO public.ppcac_conflitos_pe (
                    nome_area, municipio, municipio_ibge, proprietario, movimento_social,
                    processos_judiciais, processos_mppe, processos_sei,
                    ano_referencia, situacao, observacoes,
                    datajud_processos, despejo_zero_ids, sigef_codigos,
                    car_codigos, iterpe_nomes, incra_projetos, total_car_imoveis,
                    fonte_geometria, tipo_geometria, tem_geometria_exata,
                    centroid_lat, centroid_lon, bbox, geometry
                ) VALUES (
                    :nome_area, :municipio, :municipio_ibge, :proprietario, :movimento_social,
                    :processos_judiciais, :processos_mppe, :processos_sei,
                    :ano_referencia, :situacao, :observacoes,
                    :datajud_processos, :despejo_zero_ids, :sigef_codigos,
                    :car_codigos, :iterpe_nomes, :incra_projetos, :total_car_imoveis,
                    :fonte_geometria, :tipo_geometria, :tem_geometria_exata,
                    :centroid_lat, :centroid_lon, :bbox,
                    ST_GeomFromText(:geom_wkt, 4326)
                )
            """), {
                "nome_area": a["nome_area"],
                "municipio": mun_name,
                "municipio_ibge": municipio_ibge,
                "proprietario": a["proprietario"],
                "movimento_social": a["movimento_social"],
                "processos_judiciais": a["processos_judiciais"],
                "processos_mppe": a["processos_mppe"],
                "processos_sei": a["processos_sei"],
                "ano_referencia": a["ano_referencia"],
                "situacao": a["situacao"],
                "observacoes": obs_str,
                "datajud_processos": matched_dj_processes,
                "despejo_zero_ids": matched_dz_ids,
                "sigef_codigos": sigef_matched_codes,
                "car_codigos": car_matched_codes,
                "iterpe_nomes": iterpe_matched_names,
                "incra_projetos": incra_matched_projs,
                "total_car_imoveis": len(car_matched_codes),
                "fonte_geometria": fonte_geometria,
                "tipo_geometria": tipo_geometria,
                "tem_geometria_exata": tem_geometria_exata,
                "centroid_lat": centroid_lat,
                "centroid_lon": centroid_lon,
                "bbox": bbox,
                "geom_wkt": geom_wkt,
            })
            inserted_count += 1

        # Create indexes
        conn.execute(text("CREATE INDEX idx_ppcac_conflitos_geometry ON public.ppcac_conflitos_pe USING GIST(geometry);"))
        conn.execute(text("CREATE INDEX idx_ppcac_conflitos_municipio ON public.ppcac_conflitos_pe(municipio_ibge);"))
        conn.execute(text("CREATE INDEX idx_ppcac_conflitos_situacao ON public.ppcac_conflitos_pe(situacao);"))

        logger.info(f"Successfully ingested {inserted_count} PPCAC conflict areas into PostGIS.")

    return inserted_count


if __name__ == "__main__":
    count = ingest_ppcac()
    print(f"PPCAC ingestion complete: {count} areas processed.")
